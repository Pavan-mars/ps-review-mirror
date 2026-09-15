#!/usr/bin/env python3
"""Generate CUBIC MARS Chicago end-to-end pipeline Word guide."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "CUBIC_MARS_Chicago_End_to_End_Pipeline_Guide.docx"


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def build() -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_heading("CUBIC MARS Chicago (Ventra)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("End-to-End Data and Machine Learning Pipeline")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].bold = True
    sub.runs[0].font.size = Pt(14)
    meta = doc.add_paragraph(f"Mars Technologies | {date.today().strftime('%B %Y')}")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    para(
        doc,
        "This guide explains how Chicago Ventra fare-collection operations data is collected, refined, "
        "turned into modelling features, scored by machine learning models, and presented to maintenance "
        "and analytics teams. It is written for readers who need the full story—from raw telemetry to "
        "daily predictions—without assuming prior knowledge of the implementation repository.",
    )

    # --- 1 Business ---
    doc.add_heading("1. Why this platform exists", level=1)
    para(
        doc,
        "CTA Ventra relies on ticket vending machines (TVMs), station gates, and bus validators to process "
        "millions of taps and purchases. When a device fails, riders experience delays, revenue is lost, "
        "and contractual service levels may be breached. The MARS programme uses historical telemetry to "
        "anticipate failures, explain fault patterns, detect unusual behaviour early, and estimate how long "
        "components may last before replacement.",
    )
    para(doc, "Five analytical products share one data foundation but answer different questions:")
    bullets(
        doc,
        [
            "Failure prediction — Will this device have a hardware outage in the next few days?",
            "Cascade analysis — When one subsystem fails, which others tend to follow in sequence?",
            "Root cause — When an incident opens, what subsystem and severity best describe it?",
            "Anomaly detection — Is this hour of operation abnormal compared with the device’s own baseline?",
            "Reliability / remaining useful life — How long until this component is likely to fail again?",
        ],
    )
    para(
        doc,
        "Modelling focuses on in-service TVMs, rail gates, and bus validators. Back-office retail devices, "
        "legacy turnstiles, and pure network gear are out of scope because they do not behave like the "
        "fare-collection fleet under study.",
    )
    # --- 2 Source data ---
    doc.add_heading("2. Source data: what the devices and back-office systems send", level=1)
    para(
        doc,
        "Most data originates in Cubic’s Oracle operational data store (ODS), replicated over a secure "
        "connection into a cloud lakehouse. A smaller set comes from ServiceNow (incidents and configuration). "
        "Together these sources describe both what devices did and how operators recorded failures.",
    )
    doc.add_heading("2.1 Device telemetry and events", level=2)
    para(
        doc,
        "Device event streams are the behavioural heartbeat of the fleet. Each message typically carries a "
        "device identifier, timestamp, event type code, and severity. Events encode subsystem activity: "
        "bill and coin handling, printers, gate mechanics, card readers, communications, and system state "
        "changes. Out-of-service (OOS) events mark periods when a device cannot serve passengers. These "
        "streams power failure labels, cascade sequences, pre-incident windows, and hourly anomaly baselines.",
    )

    doc.add_heading("2.2 Tap and payment transactions", level=2)
    para(
        doc,
        "Tap tables record each card or mobile payment attempt: approval or rejection codes, fare amounts, "
        "operator, and time. Separate read-transaction feeds cover validator-style reads. Transaction volume "
        "and reject rates indicate customer-facing health—a device may log many internal warnings while still "
        "accepting taps, or may reject a rising share of taps before a full outage. Feature engineering "
        "maps raw status codes into approved vs rejected counts, excluding known non-decision codes (for "
        "example stale-tap timing metrics that are not passenger-facing rejections).",
    )

    doc.add_heading("2.3 Metrics and counters", level=2)
    para(
        doc,
        "High-frequency metric tables store rolling counters (message counts, timing statistics, health "
        "indices) keyed by device and service day. These support daily and hourly aggregates—such as average "
        "tap timing or message rates—that smooth sparse events into stable signals for anomaly detection.",
    )

    doc.add_heading("2.4 Availability, outages, and KPIs", level=2)
    para(
        doc,
        "Availability tables summarise when devices were down, for how long, and under which failure "
        "taxonomy level (for example non-payment vs full loss of function). KPI summaries compare actual "
        "performance to targets by day and facility. These feed both supervised labels (did the device fail?) "
        "and operational context (availability percentage, breach flags).",
    )

    doc.add_heading("2.5 Hardware configuration and components", level=2)
    para(
        doc,
        "Hardware configuration snapshots list installed components—serial numbers, descriptions, last "
        "reported change dates—linked to devices. This is the authoritative source for component age and "
        "identity in reliability modelling, because device-level serial fields are often incomplete compared "
        "with component-level records.",
    )

    doc.add_heading("2.6 Incidents and service management", level=2)
    para(
        doc,
        "ServiceNow (and Oracle-mirrored CTA availability extracts) provide incident text, priority, "
        "timestamps, and links to assets. Root-cause modelling uses structured fields and descriptions; "
        "incident history also enriches availability events with operational context. SaaS incident feeds "
        "must stay in sync with device telemetry or models will mix fresh events with stale tickets.",
    )

    doc.add_heading("2.7 Reference dimensions", level=2)
    para(
        doc,
        "Smaller tables decode codes: device types, facilities, stops, event types, cashbox events, calendar "
        "dates, and failure-level taxonomies agreed with the agency. Dimensions attach human-readable labels "
        "and fleet category (TVM, gate, validator) to every fact row.",
    )

    # --- 3 Medallion ---
    doc.add_heading("3. From raw landing to analytics-ready tables", level=1)
    para(
        doc,
        "Data moves through four logical layers. Each layer adds quality, consistency, and business meaning "
        "so downstream SQL and notebooks do not re-implement the same joins and rules.",
    )
    numbered(
        doc,
        [
            "Raw — Append-only files as landed from source, partitioned by load or business date, with batch identifiers for audit.",
            "Bronze — Typed tables mirroring source schemas in the lakehouse; incremental merges or appends keep them current.",
            "Silver — Cleaned grains: one row per device-day, device-hour, outage interval, or incident; derived flags (chargeable failure, subsystem, severity).",
            "Gold — Problem-specific feature tables aligned to each model’s grain and label definition.",
        ],
    )
    para(
        doc,
        "Incremental ingestion uses contract metadata: which watermark column advances, whether to merge keys "
        "or append only, and how to bound date probes so future-dated sentinel rows in source do not fake "
        "freshness. After major loads, automated checks validate primary keys, allowed code sets, and row "
        "counts against expectations.",
    )
    # --- 4 Feature engineering ---
    doc.add_heading("4. Feature engineering: purpose of major variable families", level=1)
    para(
        doc,
        "Features are deliberately split into same-day signals, rolling history, and context so models see "
        "what was knowable at prediction time—not outcomes from the future.",
    )

    doc.add_heading("4.1 Shared building blocks", level=2)
    bullets(
        doc,
        [
            "Subsystem event counts (bill handling, coin handling, printer, gate mechanics, reader, communications) — localize which part of the device is stressed.",
            "Critical and OOS counts — severity-weighted stress and periods when the device cannot serve riders.",
            "Rolling 7-day and 30-day aggregates — short memory of recent behaviour; smooths day-to-day noise.",
            "Tap volume and reject rate — passenger-visible performance; often leads internal fault logs.",
            "Availability and outage duration — SLA-oriented stress; chargeable outage flags restrict labels to hardware-relevant failures.",
            "Device category and facility — control for different mechanics (TVM vs gate vs bus validator) and station workload.",
        ],
    )

    doc.add_heading("4.2 Failure prediction features", level=2)
    para(
        doc,
        "Grain: one row per device per service day. Labels indicate hardware out-of-service within a forward "
        "window (for example three days). Same-day counters capture acute stress; rolling means (three- and "
        "seven-day) capture deteriorating trends. Engineered ratios—such as event density relative to "
        "availability, or severity combined with low availability—highlight devices that are busy yet unstable. "
        "Category encoding ensures tree models split appropriately across fleet types without treating IDs as numeric.",
    )

    doc.add_heading("4.3 Cascade analysis features", level=2)
    para(
        doc,
        "Grain: one row per multi-event chain per day. Sequences are stored as ordered subsystem paths "
        "(for example coin handling leading to communications then gate fault). Features describe chain "
        "length, duration, presence of cash-handling stages, and whether critical or OOS severities appear "
        "mid-chain. Purpose: maintenance can intervene after early links instead of waiting for full escalation.",
    )

    doc.add_heading("4.4 Root cause features", level=2)
    para(
        doc,
        "Grain: one row per availability or incident event. Training uses only information available before "
        "or at incident start: event counts in the prior 24 hours and 7 days, broken out by subsystem. Text "
        "fields from fault descriptions support NLP models. Targets include severity level under the Ventra "
        "taxonomy and a coarse root-cause category (bill handling, coin handling, gate, reader, communications, "
        "power, printer, other). Pre-incident windows are strict to avoid leaking the incident outcome into features.",
    )

    doc.add_heading("4.5 Anomaly detection features", level=2)
    para(
        doc,
        "Grain: one row per device per hour. Three parallel signals are computed: (1) hourly event rate vs "
        "a device-specific rolling mean and standard deviation; (2) key metric counters vs baseline; "
        "(3) daily tap reject rate against a fixed operational threshold. Cyclical time encodings (hour and "
        "day of week as sine/cosine) capture rush-hour effects without treating hour 23 as far from hour 0. "
        "An hour is flagged anomalous when at least two signals fire, reducing false alarms from a single noisy channel.",
    )

    doc.add_heading("4.6 Reliability and component life features", level=2)
    para(
        doc,
        "Grain: one row per component serial on a device. Age since install or last change, cumulative failure "
        "count, mean time between failures, and cashbox-related activity provide covariates for survival "
        "models. Censoring is explicit: many components have not failed yet, so their “time to failure” is "
        "unknown but age is informative. Risk scores normalize failure counts by age to compare young vs old parts fairly.",
    )

    # --- 5 Algorithms ---
    doc.add_heading("5. Algorithms used and why they were chosen", level=1)

    doc.add_heading("5.1 Failure prediction", level=2)
    para(
        doc,
        "Problem type: binary classification with heavy class imbalance and time-ordered rows. Gradient-boosted "
        "trees (LightGBM, XGBoost, CatBoost) capture nonlinear interactions among subsystem counters. A random "
        "forest adds diversity; logistic regression provides a linear baseline. A stacked ensemble combines "
        "base model probabilities with a meta learner, which often improves precision–recall on rare positive "
        "days. Training uses time-series cross-validation (no random shuffle) so future days never leak into "
        "past training folds. Primary evaluation metrics are ROC-AUC and average precision; deployment thresholds "
        "trade false alarms against missed failures. Production direction is batch scoring on a daily feature "
        "file rather than always-on real-time endpoints, for cost and reproducibility.",
    )
    doc.add_heading("5.2 Cascade analysis", level=2)
    para(
        doc,
        "Problem type: unsupervised structure discovery on sequences. Methods include frequent sequence mining, "
        "Markov transition estimates between subsystems, network centrality on co-occurrence graphs, and "
        "clustering of chain shapes. No single binary label exists; the output is interpretable patterns "
        "(for example cashbox jams followed by communication faults) used in dashboards and playbooks rather "
        "than a single probability score.",
    )

    doc.add_heading("5.3 Root cause classification", level=2)
    para(
        doc,
        "Problem type: multiclass classification with many sparse severity levels. Gradient boosting and "
        "regularized linear models handle mixed numeric pre-incident counts; text fields may use TF-IDF or "
        "similar bag-of-words features for description mining. Rare classes are merged or capped when support "
        "is too small to learn stably. Macro-F1 across classes is preferred over accuracy because incidents "
        "are imbalanced across severity codes.",
    )

    doc.add_heading("5.4 Anomaly detection", level=2)
    para(
        doc,
        "Problem type: extremely rare positive hours. A rule-based ensemble (two of three statistical signals) "
        "gives transparent operations logic. Unsupervised isolation forest or similar methods can explore "
        "multivariate outliers offline. Per-device adaptive baselines matter because a busy downtown gate "
        "naturally has higher event rates than a quiet station—global thresholds would over-alert.",
    )
    doc.add_heading("5.5 Survival and remaining useful life", level=2)
    para(
        doc,
        "Problem type: regression with censoring. Kaplan–Meier curves estimate survival over time without "
        "covariates. Cox proportional hazards and Weibull accelerated-failure-time models incorporate "
        "component type, age, and failure history while treating un-failed parts as censored observations—"
        "essential when only a small fraction have seen a second failure. Outputs include survival probability "
        "curves, expected days to failure, and fleet rankings (concordance index) for prioritizing replacements.",
    )

    # --- 6 Production flow ---
    doc.add_heading("6. Production flow: from daily ingest to dashboard", level=1)
    numbered(
        doc,
        [
            "Ingest new and changed rows from Oracle (and ServiceNow where applicable) into raw storage, then merge into bronze tables.",
            "Rebuild or incrementally refresh silver tables in dependency order (dimensions before facts, events before outages).",
            "Build gold feature tables for each problem statement; run data-quality checks.",
            "Export scored-ready datasets to object storage if training or batch jobs run outside the lakehouse.",
            "Train or apply models: daily batch scoring for failure prediction and incidents; scheduled jobs for cascades, anomalies, and weekly reliability refits.",
            "Write prediction outputs and analytics tables to object storage in agreed formats (columnar files preferred).",
            "Loader functions copy artifacts into a relational operations database with upsert semantics so reruns are idempotent.",
            "API layer exposes read routes per problem statement; the web dashboard queries the API only (never object storage directly).",
        ],
    )
    para(
        doc,
        "Orchestration is intended to be event-driven: when the gold layer successfully completes for a service "
        "date, a cloud event triggers scoring and load steps. Until that chain is fully automated, some steps "
        "may run on a schedule or manually while still following the same logical order.",
    )
    # --- 7 MLOps and drift ---
    doc.add_heading("7. MLOps, monitoring, and drift", level=1)
    doc.add_heading("7.1 Experiment tracking and model registry", level=2)
    para(
        doc,
        "Training runs record parameters, metrics, and artifact locations in an experiment tracking service. "
        "Approved models are registered with version numbers and promoted only after offline gates pass. "
        "Feature column order and types are frozen in contract documents so batch scorers match training exactly.",
    )
    doc.add_heading("7.2 Data drift", level=2)
    para(
        doc,
        "Schema checks and distribution comparisons on incoming bronze and silver columns detect source changes "
        "(new codes, missing columns, shifted date ranges). Continuity audits look for gaps in service days "
        "along the event spine. These catch upstream feed stops before models silently score empty frames.",
    )
    doc.add_heading("7.3 Model and concept drift", level=2)
    para(
        doc,
        "After deployment, score distributions and key input features are compared to training baselines. "
        "Sudden shifts in predicted failure rate or anomaly rate may indicate firmware changes, fleet mix "
        "changes, or broken features. Endpoint or batch monitors can schedule reports; operations dashboards "
        "should show data-as-of dates so stale scores are visible.",
    )
    doc.add_heading("7.4 Observability", level=2)
    para(
        doc,
        "Cloud logging for serverless loaders and APIs is the first line of troubleshooting. Metrics cover "
        "job success, latency, and invocation counts. Broader deployments may add centralized dashboards "
        "(for example Grafana) and metric stores (Prometheus or cloud-native equivalents); the architectural "
        "goal is the same: alert when ingestion, scoring, or load fails, not when a user notices old numbers on screen.",
    )
    # --- 8 CI/CD ---
    doc.add_heading("8. How build, test, and deployment (CI/CD) works", level=1)
    para(
        doc,
        "CI/CD separates what runs on every code change from what runs when promoting a release to an environment.",
    )
    doc.add_heading("8.1 Continuous integration (on each change)", level=2)
    para(
        doc,
        "When developers push changes or open a merge request, automated jobs typically verify commit message "
        "policy, run static checks on SQL and application code, execute unit tests on loader and scoring logic, "
        "and build packages (for example zipped serverless functions or container images) to prove they compile. "
        "Pull requests should not merge if checks fail. Data pipeline SQL may be validated for syntax and "
        "breaking grain assumptions without running full cluster jobs in CI.",
    )
    doc.add_heading("8.2 Continuous delivery (promotion to dev / test)", level=2)
    para(
        doc,
        "Approved mainline commits trigger deployment to a non-production account: update serverless function "
        "code with version aliases, push container images to a registry with vulnerability scanning, apply "
        "infrastructure templates if needed, and run smoke tests (health endpoints, sample API queries, loader "
        "dry-run mode that reads object storage without writing). Database schema migrations run in a controlled "
        "order with idempotent scripts so re-deploy does not wipe scored tables.",
    )
    doc.add_heading("8.3 Production release", level=2)
    para(
        doc,
        "Production promotion adds manual approval, pinned image digests or function versions, and rollback "
        "instructions (revert alias to previous version). Data jobs—lakehouse incremental load, silver/gold "
        "rebuild—are released on compatible schedules with application code so feature columns and model "
        "contracts stay aligned. Machine learning releases tie model registry version to the git tag that "
        "produced the training artifact.",
    )
    doc.add_heading("8.4 What is automated today vs the target state", level=2)
    para(
        doc,
        "Governance automation (commit policy on push) may exist while full build-and-deploy pipelines are still "
        "maturing. Many components are deployed via scripted packages from engineering workstations or cloud "
        "shell environments. The target is a single visible pipeline: code merge → build artifacts → deploy to "
        "dev → automated smoke tests → approval → production, for application code, infrastructure, and "
        "scheduled data jobs alike.",
    )
    # --- 9 Challenges ---
    doc.add_heading("9. Challenges encountered and how they were addressed", level=1)
    challenges = [
        (
            "Source connectivity and credentials",
            "VPN MTU issues and database password rotation required JDBC tuning, bounded queries, and secret store updates.",
        ),
        (
            "Future-dated and sentinel rows in source",
            "Watermark probes and silver filters exclude impossible service dates so freshness checks stay honest.",
        ),
        (
            "Device vs component identity",
            "Modelling uses stable device keys and component serial tables rather than incomplete device serial fields or mismatched CMDB bridges.",
        ),
        (
            "Label definition",
            "Hardware chargeable failures use agreed failure-level taxonomy; validator failures are derived from event streams where availability summaries are sparse.",
        ),
        (
            "Tap approval codes",
            "Domain validation on status codes separates real rejections from timing baselines and override codes.",
        ),
        (
            "Daily vs one-time model outputs",
            "Batch daily scoring and gold-complete triggers replace re-loading static test-set exports so dashboards reflect new service days.",
        ),
        (
            "Mixed freshness (telemetry vs incidents)",
            "Incident feeds require their own incremental path so root-cause views do not lag device events by months.",
        ),
    ]
    for head, body in challenges:
        p = doc.add_paragraph()
        p.add_run(head + ". ").bold = True
        p.add_run(body)

    doc.add_heading("10. Summary", level=1)
    para(
        doc,
        "The Chicago Ventra MARS pipeline turns heterogeneous operational data—events, taps, metrics, outages, "
        "configuration, and tickets—into layered tables, problem-specific features, and scored outputs consumed "
        "by maintenance analytics. Tree ensembles, sequence analytics, multiclass classifiers, rule ensembles, "
        "and survival models were chosen to match each question’s grain, label rarity, and need for interpretability. "
        "Reliable production depends equally on ingestion contracts, quality gates, drift awareness, and a clear "
        "CI/CD path from code change to deployed loaders, APIs, and scheduled scores.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
